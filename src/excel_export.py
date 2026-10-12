"""Excel summary workbook for the dislocation analysis (openpyxl).

Produces a formatted, Excel-ready workbook; the same tables are also
available as CSVs. ALL DATA IS SYNTHETIC.
"""

from __future__ import annotations

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

TITLE_FILL = PatternFill("solid", fgColor="1F3864")
TITLE_FONT = Font(color="FFFFFF", bold=True, size=12)
HEADER_FILL = PatternFill("solid", fgColor="D9E2F3")
HEADER_FONT = Font(bold=True)
THIN = Side(style="thin", color="B0B0B0")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def _style_sheet(ws, title: str, headers: list[str], rows: list[tuple]) -> None:
    ws["A1"] = title
    ws["A1"].font = TITLE_FONT
    ws["A1"].fill = TITLE_FILL
    ws.merge_cells(start_row=1, start_column=1,
                   end_row=1, end_column=len(headers))
    ws["A2"] = "SYNTHETIC DATA - generated for demonstration"
    ws["A2"].font = Font(italic=True, color="808080", size=9)
    for c, h in enumerate(headers, start=1):
        cell = ws.cell(row=4, column=c, value=h)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.border = BORDER
    for r, row in enumerate(rows, start=5):
        for c, value in enumerate(row, start=1):
            cell = ws.cell(row=r, column=c, value=value)
            cell.border = BORDER
            if isinstance(value, float):
                cell.number_format = "0.0000" if abs(value) < 10 else "#,##0.00"
    for c in range(1, len(headers) + 1):
        width = max(len(str(ws.cell(row=r, column=c).value or ""))
                       for r in range(4, 5 + len(rows)))
        ws.column_dimensions[get_column_letter(c)].width = min(max(width + 2, 14), 34)
    ws.sheet_properties.pageSetUpPr = None


def _frame_rows(frame, columns):
    return [tuple(r) for r in frame[columns].itertuples(index=False, name=None)]


def build_workbook(results: dict, path: str) -> None:
    """Build the dislocation summary workbook from SQL result frames."""
    wb = Workbook()

    ws = wb.active
    ws.title = "Cover"
    ws["A1"] = "Rate-Change Dislocation Analysis - Summary"
    ws["A1"].font = Font(bold=True, size=14, color="1F3864")
    ws["A3"] = "SYNTHETIC DATA - all policies, premiums and rate tables are generated."
    ws["A3"].font = Font(italic=True, color="808080")
    ws["A5"] = "Pricing versions compared: v1 (in force) vs v2 (candidate)"
    ws["A6"] = f"Policies analyzed: {results['n_policies']:,}"
    ws["A7"] = "Renewal guardrails: increases capped at +15%, decreases floored at -10%"
    ws["A9"] = "Sheets: Overall | By_State | By_Breed_Group | By_Age_Band | Capping_Impact | Cap_Hit_By_State"
    for row in ("A1", "A5", "A6", "A7", "A9"):
        ws.column_dimensions["A"].width = 80

    overall = results["overall"].iloc[0]
    _style_sheet(
        wb.create_sheet("Overall"),
        "Overall dislocation: v1 -> v2",
        ["Metric", "Value"],
        [("Policies analyzed", int(overall["n_policies"])),
         ("Min premium change", float(overall["min_change_pct"])),
         ("Avg premium change", float(overall["avg_change_pct"])),
         ("Max premium change", float(overall["max_change_pct"])),
         ("Avg v1 premium ($)", float(overall["avg_v1_premium"])),
         ("Avg v2 premium ($)", float(overall["avg_v2_premium"])),
         ("Total premium change ($)", float(overall["total_premium_change"]))],
    )

    _style_sheet(
        wb.create_sheet("By_State"),
        "Premium change by state (min / avg / max)",
        ["State", "Policies", "Min change", "Avg change", "Max change"],
        _frame_rows(results["by_state"],
                    ["state", "n_policies", "min_change_pct",
                     "avg_change_pct", "max_change_pct"]),
    )
    _style_sheet(
        wb.create_sheet("By_Breed_Group"),
        "Premium change by breed group (min / avg / max)",
        ["Breed group", "Policies", "Min change", "Avg change", "Max change"],
        _frame_rows(results["by_breed"],
                    ["breed_group", "n_policies", "min_change_pct",
                     "avg_change_pct", "max_change_pct"]),
    )
    _style_sheet(
        wb.create_sheet("By_Age_Band"),
        "Premium change by age band (min / avg / max)",
        ["Age band", "Policies", "Min change", "Avg change", "Max change"],
        _frame_rows(results["by_age_band"],
                    ["age_band", "n_policies", "min_change_pct",
                     "avg_change_pct", "max_change_pct"]),
    )

    cap = results["capping"].iloc[0]
    _style_sheet(
        wb.create_sheet("Capping_Impact"),
        "Renewal capping rules: +15% increase cap, -10% decrease floor",
        ["Metric", "Value"],
        [("Policies analyzed", int(cap["n_policies"])),
         ("Policies hitting increase cap", int(cap["capped_increases"])),
         ("Policies hitting decrease floor", int(cap["capped_decreases"])),
         ("Avg uncapped change", float(cap["avg_uncapped_change_pct"])),
         ("Avg capped change", float(cap["avg_capped_change_pct"])),
         ("Total uncapped premium ($)", float(cap["total_uncapped_premium"])),
         ("Total capped premium ($)", float(cap["total_capped_premium"])),
         ("Premium impact of capping ($)",
          float(cap["total_capped_premium"]) - float(cap["total_uncapped_premium"]))],
    )
    _style_sheet(
        wb.create_sheet("Cap_Hit_By_State"),
        "Cap hit rate by state",
        ["State", "Policies", "Capped policies", "Cap hit rate"],
        _frame_rows(results["cap_hit_by_state"],
                    ["state", "n_policies", "capped_policies", "cap_hit_rate"]),
    )

    ws_notes = wb.create_sheet("Notes")
    notes = [
        "Method notes",
        "1. 120,000 synthetic policies were rated under clean v1 and clean v2 tables.",
        "2. Premium change = (v2 - v1) / v1 per policy.",
        "3. Aggregations (min/max/avg) were computed in SQLite; see sql/dislocation_analysis.sql.",
        "4. Capping re-computes each renewal premium with the guardrails applied",
        "   before any further aggregation.",
        "5. ALL DATA IS SYNTHETIC and generated for demonstration purposes.",
    ]
    for i, line in enumerate(notes, start=1):
        ws_notes.cell(row=i, column=1, value=line)
        if i == 1:
            ws_notes.cell(row=i, column=1).font = Font(bold=True, size=12)
    ws_notes.column_dimensions["A"].width = 90

    wb.save(path)
