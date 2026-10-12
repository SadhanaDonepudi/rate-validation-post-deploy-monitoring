"""Pytest suite for the pre-deployment rate validation & post-deploy monitoring repo.

Every count asserted here is produced by the repo's own code on synthetic
data (fixed seeds); nothing is hand-computed.

Covers:
  * 50,000 quotes rated per pricing version
  * 100% premium match on clean tables (v1 and v2)
  * 23/23 planted defects caught on the defective v2 tables
  * 120,000 policies in the dislocation analysis
  * min/max/avg sanity of the dislocation aggregates
  * 3/3 seeded live-rating anomalies detected within 24 simulated hours
"""

import json
import os
import sys

import pytest

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, os.path.join(ROOT, "src"))

from data_gen import generate_quotes  # noqa: E402
from defect_injection import EXPECTED_DEFECT_COUNT, build_all  # noqa: E402
from dislocation import run_dislocation  # noqa: E402
from monitor import run_monitor, simulate_events  # noqa: E402
from rating_engine import RatingEngine  # noqa: E402
from validate_rates import attribute_defects, run_validation_suite  # noqa: E402

TABLES = os.path.join(ROOT, "data", "rate_tables")
OUT = os.path.join(ROOT, "outputs")
SQL = os.path.join(ROOT, "sql", "dislocation_analysis.sql")


@pytest.fixture(scope="module")
def tables_and_manifest(tmp_path_factory):
    build_all(TABLES, os.path.join(OUT, "defect_manifest.json"))
    with open(os.path.join(OUT, "defect_manifest.json"), encoding="utf-8") as fh:
        manifest = json.load(fh)["defects"]
    return manifest


@pytest.fixture(scope="module")
def quotes():
    return generate_quotes(50_000)


@pytest.fixture(scope="module")
def clean_reports(quotes, tables_and_manifest):
    reports = {}
    for version in ("v1", "v2"):
        engine = RatingEngine(os.path.join(TABLES, "clean", version), version)
        reports[version] = run_validation_suite(engine, version, quotes)
    return reports


@pytest.fixture(scope="module")
def defective_report(quotes, tables_and_manifest):
    engine = RatingEngine(os.path.join(TABLES, "defective", "v2"), "v2")
    return run_validation_suite(engine, "v2", quotes)


@pytest.fixture(scope="module")
def dislocation_results():
    return run_dislocation(os.path.join(TABLES, "clean"), SQL, 120_000)


@pytest.fixture(scope="module")
def monitoring_result():
    events, zip_map = simulate_events(os.path.join(TABLES, "clean"))
    return run_monitor(events, zip_map)


# ---- validation suite ---------------------------------------------------------

def test_quote_count_per_version(quotes, clean_reports):
    assert len(quotes) == 50_000
    assert clean_reports["v1"]["quotes_tested"] == 50_000
    assert clean_reports["v2"]["quotes_tested"] == 50_000


def test_clean_tables_100_percent_match(clean_reports):
    for version in ("v1", "v2"):
        report = clean_reports[version]
        assert report["mismatches"] == 0, f"{version}: {report['mismatches']} mismatches"
        assert report["match_rate"] == 1.0


def test_all_23_planted_defects_caught(defective_report, tables_and_manifest):
    assert EXPECTED_DEFECT_COUNT == 23
    assert len(tables_and_manifest) == 23
    attribution = attribute_defects(defective_report, tables_and_manifest)
    caught = [d for d in attribution if d["caught"]]
    missing = [d["id"] for d in attribution if not d["caught"]]
    assert len(caught) == 23, f"uncaught defects: {missing}"
    # every defect id from DEF-01..DEF-23 is present and caught
    assert {d["id"] for d in caught} == {f"DEF-{i:02d}" for i in range(1, 24)}


def test_defective_run_still_rates_all_quotes(defective_report):
    assert defective_report["quotes_tested"] == 50_000


# ---- dislocation analysis ------------------------------------------------------

def test_dislocation_policy_count(dislocation_results):
    assert dislocation_results["n_policies"] == 120_000
    assert len(dislocation_results["policies"]) == 120_000


def test_dislocation_min_max_avg_sanity(dislocation_results):
    for name in ("by_state", "by_breed", "by_age_band"):
        frame = dislocation_results[name]
        assert frame[["min_change_pct", "avg_change_pct", "max_change_pct"]].notna().all().all()
        assert ((frame["min_change_pct"] <= frame["avg_change_pct"])
                & (frame["avg_change_pct"] <= frame["max_change_pct"])).all(), name
    overall = dislocation_results["overall"].iloc[0]
    assert overall["min_change_pct"] <= overall["avg_change_pct"] <= overall["max_change_pct"]
    # sanity bounds: no policy moves more than +/-50% in this synthetic book
    assert -0.50 < overall["min_change_pct"] < 0.50
    assert -0.50 < overall["max_change_pct"] < 0.50
    assert -0.50 < overall["avg_change_pct"] < 0.50


def test_capping_bounds_respected(dislocation_results):
    cap = dislocation_results["capping"].iloc[0]
    assert -0.10 <= cap["avg_capped_change_pct"] <= 0.15
    assert cap["capped_increases"] + cap["capped_decreases"] <= cap["n_policies"]


# ---- post-deployment monitoring --------------------------------------------------

def test_all_anomalies_detected_within_24h(monitoring_result):
    assert monitoring_result["anomalies_seeded"] == 3
    assert monitoring_result["anomalies_detected"] == 3
    kinds = {a["kind"] for a in monitoring_result["alerts"]}
    assert kinds == {"RATE_DRIFT", "STALE_FACTOR_USAGE", "TERRITORY_MISMAPPING"}
    for alert in monitoring_result["alerts"]:
        assert alert["detection_latency_hours"] < 24, alert
    assert monitoring_result["all_within_24h"]


def test_monitoring_event_volume(monitoring_result):
    assert monitoring_result["n_events"] == 72 * 30
    assert monitoring_result["hours_simulated"] == 72


# ---- workbook ---------------------------------------------------------------------

def test_excel_workbook_exists():
    pytest.importorskip("openpyxl")
    path = os.path.join(OUT, "rate_change_dislocation_summary.xlsx")
    assert os.path.exists(path), "run run_all.py first to build outputs"
    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True)
    for sheet in ("Cover", "Overall", "By_State", "By_Breed_Group",
                  "By_Age_Band", "Capping_Impact", "Cap_Hit_By_State", "Notes"):
        assert sheet in wb.sheetnames, sheet
